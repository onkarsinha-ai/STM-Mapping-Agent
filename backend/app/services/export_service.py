from io import BytesIO
from typing import List
from sqlalchemy import select
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from app.database import AsyncSessionLocal
from app.models.mapping import Mapping, MappingStatus


class ExportService:
    @staticmethod
    async def generate_excel(project_id: str) -> bytes:
        async with AsyncSessionLocal() as session:
            result = await session.execute(
                select(Mapping).where(
                    Mapping.project_id == project_id,
                    Mapping.status.in_([MappingStatus.approved, MappingStatus.modified, MappingStatus.rejected])
                )
            )
            mappings = result.scalars().all()

        wb = Workbook()

        # Group by target table
        tables = {}
        for m in mappings:
            table = m.target_table
            if table not in tables:
                tables[table] = []
            tables[table].append(m)

        first = True
        for table_name, table_mappings in tables.items():
            if first:
                ws = wb.active
                ws.title = table_name
                first = False
            else:
                ws = wb.create_sheet(title=table_name)

            headers = ["Target DB", "Target Table", "Target Column", "Source DB", "Source Table",
                      "Source Column", "Business Logic", "Transformation", "Confidence Score"]

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

            for row, m in enumerate(table_mappings, 2):
                is_rejected = m.status == MappingStatus.rejected
                ws.cell(row=row, column=1, value="target")
                ws.cell(row=row, column=2, value=m.target_table)
                ws.cell(row=row, column=3, value=m.target_column)
                ws.cell(row=row, column=4, value="source")
                ws.cell(row=row, column=5, value="USER INPUT NEEDED" if is_rejected else m.source_table)
                ws.cell(row=row, column=6, value="USER INPUT NEEDED" if is_rejected else m.source_column)
                ws.cell(row=row, column=7, value="USER INPUT NEEDED" if is_rejected else m.business_logic)
                ws.cell(row=row, column=8, value="USER INPUT NEEDED" if is_rejected else m.transformation_rule)
                ws.cell(row=row, column=9, value=0.00 if is_rejected else (float(m.confidence_score) if m.confidence_score else None))

            for col in range(1, len(headers) + 1):
                ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 20

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
